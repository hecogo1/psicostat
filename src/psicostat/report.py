"""APA 7 style methods and results text (es/en) + Markdown/DOCX export."""
from __future__ import annotations

import os

import numpy as np
import pandas as pd

from . import __version__
from .design import assumption_label, recommend
from .effects import fmt_num, fmt_p, magnitude
from .i18n import t
from .power import explain, sample_size

DESIGN_TXT = {
    "between": ("entre sujetos", "between-subjects"),
    "within": ("intra sujetos (medidas repetidas)", "within-subjects (repeated measures)"),
    "mixed": ("mixto", "mixed"),
}


def _L(lang):
    return 0 if lang == "es" else 1


def stat_str(row, lang="es"):
    """APA-formatted statistic string, e.g. t(37.62) = 2.31, p = .026, g = 0.72, 95% CI [0.08, 1.36]."""
    sym = row["stat_symbol"]
    if sym == "OR":
        s = f"OR = {fmt_num(row['stat'])}"
    elif sym == "binomial":
        s = f"{'binomial exacta' if lang == 'es' else 'exact binomial'}"
    else:
        df = f"({row['df']})" if row["df"] else ""
        bounded = sym in ("r", "r_s")
        s = f"{sym}{df} = {fmt_num(row['stat'], bounded)}"
    s += f", {fmt_p(row['p'])}"
    if row["es_symbol"] and row["es"] is not None:
        s += f", {row['es_symbol']} = {fmt_num(row['es'], row['bounded'])}" if row["es_symbol"] != row["stat_symbol"] else ""
        if row["ci"] is not None and all(np.isfinite(row["ci"])):
            lab = "IC 95%" if lang == "es" else "95% CI"
            s += f", {lab} [{fmt_num(row['ci'][0], row['bounded'])}, {fmt_num(row['ci'][1], row['bounded'])}]"
    return s


def _sig(row, alpha, lang):
    if row["p"] is None:
        return ""
    if lang == "es":
        return "fue estadísticamente significativo" if row["p"] < alpha else "no fue estadísticamente significativo"
    return "was statistically significant" if row["p"] < alpha else "was not statistically significant"


def _subject(r, lang):
    lab = r["label"]
    es = lang == "es"
    if lab == "Modelo/Model":
        return "El modelo global" if es else "The overall model"
    if " vs " in lab:
        a, b = lab.split(" vs ", 1)
        return f"La diferencia entre {a} y {b}" if es else f"The difference between {a} and {b}"
    if " – " in lab:
        a, b = lab.split(" – ", 1)
        return f"La correlación entre {a} y {b}" if es else f"The correlation between {a} and {b}"
    if " × " in lab and r["stat_symbol"] in ("χ²", "OR"):
        a, b = lab.split(" × ", 1)
        return f"La asociación entre {a} y {b}" if es else f"The association between {a} and {b}"
    if r["stat_symbol"] in ("t", "z"):
        return f"El predictor {lab}" if es else f"The predictor {lab}"
    if " × " in lab:
        return f"La interacción {lab}" if es else f"The {lab} interaction"
    return f"El efecto de {lab}" if es else f"The effect of {lab}"


def _desc_sentence(res, lang):
    dsc = res.descriptives
    if dsc is None or "M" not in dsc.columns:
        return ""
    parts = []
    ordinal = res.test in ("mannwhitney", "wilcoxon", "kruskal", "friedman")
    sd = "DT" if lang == "es" else "SD"
    for _, r in dsc.iterrows():
        if ordinal:
            iqr = "RIC" if lang == "es" else "IQR"
            parts.append(f"{r['grupo/group']} (Mdn = {r['Mdn']:.2f}, {iqr} = {r['RIC/IQR']:.2f}, n = {int(r['n'])})")
        else:
            parts.append(f"{r['grupo/group']} (M = {r['M']:.2f}, {sd} = {r['DT/SD']:.2f}, n = {int(r['n'])})")
    head = "Los estadísticos descriptivos fueron: " if lang == "es" else "Descriptive statistics were: "
    return head + "; ".join(parts) + "."


def methods_text(design, lang="es"):
    d = design
    rec = recommend(d)
    ss = sample_size(d, rec)
    P = []
    if lang == "es":
        P.append(("Participantes", explain(ss, lang) +
                  " El tamaño muestral debe justificarse antes de recoger los datos (Lakens, 2022)."))
        if d.analysis == "compare":
            ds = f"Se empleó un diseño {DESIGN_TXT[d.design][0]}"
            if d.design == "mixed":
                ds += (f" {d.n_groups} × {d.n_levels_b}, con {d.iv_name} ({', '.join(d.levels)}) como factor entre sujetos "
                       f"y {d.factor_b_name} ({', '.join(d.levels_b)}) como factor intra sujetos")
            elif d.n_factors == 2:
                ds += (f" factorial {d.n_groups} × {d.n_levels_b}, con los factores {d.iv_name} ({', '.join(d.levels)}) "
                       f"y {d.factor_b_name} ({', '.join(d.levels_b)})")
            else:
                ds += f" con {d.n_groups} {'condiciones' if d.design == 'within' else 'grupos'} ({', '.join(d.levels)}) de la variable independiente {d.iv_name}"
            ds += f". La variable dependiente fue {d.dv_name}"
            if d.covariates:
                ds += f", controlando las covariables {', '.join(d.covariates)}"
            P.append(("Diseño", ds + "."))
        elif d.analysis == "correlation":
            P.append(("Diseño", f"Se empleó un diseño correlacional para examinar la relación entre {d.iv_name} y {d.dv_name}."))
        elif d.analysis == "regression":
            P.append(("Diseño", f"Se empleó un diseño correlacional-predictivo con {d.dv_name} como variable criterio y "
                                f"{', '.join(d.predictors)} como predictor(es)."))
        else:
            P.append(("Diseño", f"Se examinó la asociación entre las variables categóricas {d.iv_name} y {d.dv_name}."))
        ana = (f"El análisis principal planificado fue: {t(rec.test, lang)}. "
               f"Antes del análisis se comprobaron los supuestos: {', '.join(assumption_label(a, lang)[0].lower() + assumption_label(a, lang)[1:] for a in rec.assumptions)}. ")
        if rec.fallback:
            ana += f"Si no se cumplían los supuestos, se planificó como alternativa: {t(rec.fallback, lang)}. "
        ana += " ".join(rec.notes_text(lang)) + " "
        ana += (f"Se informan tamaños del efecto con intervalos de confianza al 95%. El nivel de significación fue α = {fmt_num(d.alpha, True)} "
                f"(bilateral). Los análisis se realizaron con PsicoStat v{__version__} (Python; SciPy y statsmodels).")
        P.append(("Análisis de datos", ana))
    else:
        P.append(("Participants", explain(ss, lang) + " Sample size should be justified before data collection (Lakens, 2022)."))
        if d.analysis == "compare":
            ds = f"A {DESIGN_TXT[d.design][1]} design was used"
            if d.design == "mixed":
                ds = (f"A {d.n_groups} × {d.n_levels_b} mixed design was used, with {d.iv_name} ({', '.join(d.levels)}) as "
                      f"between-subjects factor and {d.factor_b_name} ({', '.join(d.levels_b)}) as within-subjects factor")
            elif d.n_factors == 2:
                ds = (f"A {d.n_groups} × {d.n_levels_b} {DESIGN_TXT[d.design][1]} factorial design was used, with factors "
                      f"{d.iv_name} ({', '.join(d.levels)}) and {d.factor_b_name} ({', '.join(d.levels_b)})")
            else:
                ds += f" with {d.n_groups} {'conditions' if d.design == 'within' else 'groups'} ({', '.join(d.levels)}) of the independent variable {d.iv_name}"
            ds += f". The dependent variable was {d.dv_name}"
            if d.covariates:
                ds += f", controlling for {', '.join(d.covariates)}"
            P.append(("Design", ds + "."))
        elif d.analysis == "correlation":
            P.append(("Design", f"A correlational design was used to examine the relationship between {d.iv_name} and {d.dv_name}."))
        elif d.analysis == "regression":
            P.append(("Design", f"A correlational-predictive design was used with {d.dv_name} as outcome and "
                                f"{', '.join(d.predictors)} as predictor(s)."))
        else:
            P.append(("Design", f"The association between the categorical variables {d.iv_name} and {d.dv_name} was examined."))
        ana = (f"The planned main analysis was: {t(rec.test, lang)}. Assumptions were checked beforehand: "
               f"{', '.join(assumption_label(a, lang)[0].lower() + assumption_label(a, lang)[1:] for a in rec.assumptions)}. ")
        if rec.fallback:
            ana += f"If assumptions were not met, the planned alternative was: {t(rec.fallback, lang)}. "
        ana += " ".join(rec.notes_text(lang)) + " "
        ana += (f"Effect sizes are reported with 95% confidence intervals. The significance level was α = {fmt_num(d.alpha, True)} "
                f"(two-sided). Analyses were run with PsicoStat v{__version__} (Python; SciPy and statsmodels).")
        P.append(("Data analysis", ana))
    return P


def results_text(res, lang="es"):
    d = res.design
    es = lang == "es"
    out = []
    intro = []
    if res.n_excluded:
        intro.append(f"Se excluyeron {res.n_excluded} casos con datos incompletos; la muestra analizada fue N = {res.n}." if es else
                     f"{res.n_excluded} cases with incomplete data were excluded; the analysed sample was N = {res.n}.")
    else:
        intro.append(f"La muestra analizada fue N = {res.n}." if es else f"The analysed sample was N = {res.n}.")
    ds = _desc_sentence(res, lang)
    if ds:
        intro.append(ds)
    out.append(" ".join(intro))
    # assumptions
    if res.assumptions:
        bad = [a for a in res.assumptions if not a["ok"]]
        if bad:
            s = ("Se detectaron incumplimientos de supuestos: " if es else "Assumption violations were detected: ") + "; ".join(
                f"{assumption_label(a['key'], lang)} ({a['detail']})" for a in bad) + "."
        else:
            s = "Se cumplieron los supuestos evaluados." if es else "The assumptions checked were met."
        if res.fallback_reason:
            s += (f" Por ello se aplicó la alternativa planificada: {t(res.test, lang)} (en lugar de: {t(res.planned_test, lang)})." if es else
                  f" Therefore the planned alternative was used: {t(res.test, lang)} (instead of: {t(res.planned_test, lang)}).")
        elif res.test != res.planned_test:
            s += f" Prueba aplicada: {t(res.test, lang)}." if es else f" Test used: {t(res.test, lang)}."
        out.append(s)
    for note in res.notes:
        out.append(note[_L(lang)])
    # main results
    tl = t(res.test, lang)
    sents = []
    for r in res.rows:
        mag = magnitude(r["es_symbol"], r["es"]) if r["es_symbol"] else None
        magtxt = ""
        if mag:
            magtxt = (f"; tamaño del efecto {t(mag, lang)}" if es else f"; {t(mag, lang)} effect size")
        sents.append(f"{_subject(r, lang)} {_sig(r, d.alpha, lang)}, {stat_str(r, lang)}{magtxt}.")
    head = (f"Resultados ({tl}): " if es else f"Results ({tl}): ")
    out.append(head + " ".join(sents))
    if res.posthoc is not None and "p_ajust/p_adj" in res.posthoc.columns and "A" in res.posthoc.columns:
        sig = res.posthoc[res.posthoc["p_ajust/p_adj"] < d.alpha]
        if len(sig):
            pairs = "; ".join(f"{r['A']} vs {r['B']} ({fmt_p(r['p_ajust/p_adj'])})" for _, r in sig.iterrows())
            out.append(f"Las comparaciones post hoc ({res.posthoc_method}) mostraron diferencias significativas en: {pairs}." if es else
                        f"Post hoc comparisons ({res.posthoc_method}) showed significant differences for: {pairs}.")
        else:
            out.append(f"Ninguna comparación post hoc ({res.posthoc_method}) fue significativa." if es else
                       f"No post hoc comparison ({res.posthoc_method}) was significant.")
    out.append("Recuerda: un resultado no significativo no demuestra ausencia de efecto." if es and any(r["p"] is not None and r["p"] >= d.alpha for r in res.rows)
               else ("Remember: a non-significant result does not show the absence of an effect." if (not es and any(r["p"] is not None and r["p"] >= d.alpha for r in res.rows)) else ""))
    return [p for p in out if p]


def _tables(res, lang):
    es = lang == "es"
    tabs = []
    if res.descriptives is not None:
        tabs.append(("Estadísticos descriptivos" if es else "Descriptive statistics", res.descriptives))
    if res.assumptions:
        tabs.append(("Comprobación de supuestos" if es else "Assumption checks",
                     pd.DataFrame([{("supuesto" if es else "assumption"): assumption_label(a["key"], lang), ("detalle" if es else "detail"): a["detail"],
                                    ("cumple" if es else "met"): ("sí" if es else "yes") if a["ok"] else "no"} for a in res.assumptions])))
    for k, v in res.extra_tables.items():
        name = {"coefficients": ("Coeficientes", "Coefficients"), "model": ("Ajuste del modelo", "Model fit"),
                "adjusted_means": ("Medias ajustadas", "Adjusted means")}.get(k, (k, k))[_L(lang)]
        tabs.append((name, v))
    if res.posthoc is not None:
        tabs.append(((f"Post hoc: {res.posthoc_method}"), res.posthoc))
    return tabs


def build_report(res, lang="es"):
    return {"title": res.design.title or ("Informe de análisis" if lang == "es" else "Analysis report"),
            "question": res.design.question, "methods": methods_text(res.design, lang),
            "results": results_text(res, lang), "tables": _tables(res, lang),
            "figures": [f for f in res.figures if f], "lang": lang}


def _md_table(df: pd.DataFrame):
    df = df.copy()
    for c in df.columns:
        if pd.api.types.is_float_dtype(df[c]):
            df[c] = df[c].map(lambda x: "" if pd.isna(x) else (fmt_p(x).replace("p = ", "").replace("p ", "") if str(c).startswith("p") else f"{x:.2f}"))
    cols = [str(c).replace("|", "\\|") for c in df.columns]
    lines = ["| " + " | ".join(cols) + " |", "|" + "---|" * len(cols)]
    for _, r in df.iterrows():
        lines.append("| " + " | ".join(str(v).replace("|", "\\|") for v in r.values) + " |")
    return "\n".join(lines)


def to_markdown(rep, path=None, rel_to=None):
    es = rep["lang"] == "es"
    L = [f"# {rep['title']}", ""]
    if rep["question"]:
        L += [f"**{'Pregunta de investigación' if es else 'Research question'}:** {rep['question']}", ""]
    L += [f"## {'Método' if es else 'Method'}", ""]
    for h, p in rep["methods"]:
        L += [f"### {h}", "", p, ""]
    L += [f"## {'Resultados' if es else 'Results'}", ""]
    L += [p + "\n" for p in rep["results"]]
    for i, (name, df) in enumerate(rep["tables"], 1):
        L += [f"**{'Tabla' if es else 'Table'} {i}**", "", f"*{name}*", "", _md_table(df), ""]
    for i, f in enumerate(rep["figures"], 1):
        src = os.path.relpath(f, rel_to) if rel_to else f
        L += [f"**{'Figura' if es else 'Figure'} {i}**", "", f"![]({src})", ""]
    L += ["---", "", f"*{t('disclaimer', rep['lang'])}*", ""]
    md = "\n".join(L)
    if path:
        with open(path, "w", encoding="utf-8") as fh:
            fh.write(md)
    return md


def to_docx(rep, path):
    from docx import Document
    from docx.shared import Inches, Pt

    es = rep["lang"] == "es"
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    doc.add_heading(rep["title"], 0)
    if rep["question"]:
        doc.add_paragraph(rep["question"])
    doc.add_heading("Método" if es else "Method", 1)
    for h, p in rep["methods"]:
        doc.add_heading(h, 2)
        doc.add_paragraph(p)
    doc.add_heading("Resultados" if es else "Results", 1)
    for p in rep["results"]:
        doc.add_paragraph(p)
    for i, (name, df) in enumerate(rep["tables"], 1):
        doc.add_paragraph().add_run(f"{'Tabla' if es else 'Table'} {i}").bold = True
        doc.add_paragraph().add_run(name).italic = True
        tb = doc.add_table(rows=1, cols=len(df.columns))
        tb.style = "Table Grid"
        for j, c in enumerate(df.columns):
            tb.rows[0].cells[j].text = str(c)
        for _, r in df.iterrows():
            cells = tb.add_row().cells
            for j, v in enumerate(r.values):
                cells[j].text = (f"{v:.2f}" if isinstance(v, (float, np.floating)) and not pd.isna(v) else str(v))
        doc.add_paragraph()
    for i, f in enumerate(rep["figures"], 1):
        doc.add_paragraph().add_run(f"{'Figura' if es else 'Figure'} {i}").bold = True
        doc.add_picture(f, width=Inches(5.5))
    doc.add_paragraph().add_run(t("disclaimer", rep["lang"])).italic = True
    doc.save(path)
    return path
